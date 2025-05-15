import PyPDF2
import docx
import subprocess
import os
import re
import json
import requests
import base64 # Added for base64 encoding
from flask import current_app
from app.services.llm_service import ClaudeService, LLMAnalysisError
from fpdf import FPDF # Added for TXT to PDF conversion

class DocumentProcessingError(Exception):
    pass

class FileConverter:
    @staticmethod
    def convert_to_pdf(file_path, output_dir):
        original_filename = os.path.basename(file_path)
        filename_stem, original_ext = os.path.splitext(original_filename)
        original_ext = original_ext.lower()
        pdf_filename = f"{filename_stem}.pdf"
        pdf_output_path = os.path.join(output_dir, pdf_filename)

        if not os.path.exists(output_dir):
            os.makedirs(output_dir)

        if original_ext == ".pdf":
            current_app.logger.info(f"File {original_filename} is already a PDF.")
            if os.path.abspath(file_path) == os.path.abspath(pdf_output_path):
                current_app.logger.info(f"Source and destination for PDF {original_filename} are the same. Skipping copy.")
                return pdf_output_path
            else:
                try:
                    current_app.logger.info(f"Copying PDF {file_path} to {pdf_output_path}.")
                    subprocess.run(["cp", file_path, pdf_output_path], check=True)
                    return pdf_output_path
                except subprocess.CalledProcessError as e:
                    current_app.logger.error(f"Failed to copy PDF {file_path} to {pdf_output_path}: {e}")
                    raise DocumentProcessingError(f"Failed to copy PDF: {e}")

        elif original_ext == ".docx":
            try:
                current_app.logger.info(f"Converting DOCX {original_filename} to PDF at {pdf_output_path}.")
                subprocess.run([
                    "libreoffice", "--headless", "--convert-to", "pdf",
                    "--outdir", output_dir, file_path
                ], check=True, timeout=60)
                expected_converted_file = os.path.join(output_dir, f"{filename_stem}.pdf")
                if not os.path.exists(expected_converted_file):
                    current_app.logger.error(f"Conversion of {original_filename} to PDF failed or file not found at {expected_converted_file}.")
                    raise DocumentProcessingError(f"DOCX to PDF conversion failed for {original_filename}. Output file not found.")
                current_app.logger.info(f"Successfully converted {original_filename} to {expected_converted_file}")
                return expected_converted_file
            except subprocess.TimeoutExpired:
                current_app.logger.error(f"Timeout during DOCX to PDF conversion for {original_filename}.")
                raise DocumentProcessingError(f"Timeout converting {original_filename} to PDF.")
            except subprocess.CalledProcessError as e:
                current_app.logger.error(f"Error converting DOCX {original_filename} to PDF: {e}. Output: {e.output}, Stderr: {e.stderr}")
                raise DocumentProcessingError(f"Error converting {original_filename} to PDF: {e}")
            except Exception as e:
                current_app.logger.error(f"Unexpected error during DOCX to PDF conversion for {original_filename}: {e}", exc_info=True)
                raise DocumentProcessingError(f"Unexpected error converting {original_filename} to PDF: {e}")

        elif original_ext == ".txt":
            try:
                current_app.logger.info(f"Converting TXT {original_filename} to PDF at {pdf_output_path}.")
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font("Arial", size=12)
                with open(file_path, "r", encoding="utf-8") as f:
                    for line in f:
                        pdf.cell(200, 10, txt=line, ln=1)
                pdf.output(pdf_output_path, "F")
                current_app.logger.info(f"Successfully converted {original_filename} to {pdf_output_path}")
                return pdf_output_path
            except Exception as e:
                current_app.logger.error(f"Error converting TXT {original_filename} to PDF: {e}", exc_info=True)
                raise DocumentProcessingError(f"Error converting {original_filename} to PDF: {e}")
        else:
            raise DocumentProcessingError(f"Unsupported file type for PDF conversion: {original_ext}")

class TextExtractor: 
    @staticmethod
    def extract_text_from_pdf_pypdf2(pdf_path):
        current_app.logger.info(f"Attempting to extract text from PDF (PyPDF2): {pdf_path}")
        try:
            with open(pdf_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                text = ""
                for page_num in range(len(reader.pages)):
                    page = reader.pages[page_num]
                    text += page.extract_text() or "" 
                current_app.logger.info(f"Successfully extracted text from PDF (PyPDF2): {pdf_path}. Length: {len(text)}")
                return text
        except Exception as e:
            current_app.logger.error(f"Error extracting text from PDF (PyPDF2) {pdf_path}: {e}", exc_info=True)
            raise DocumentProcessingError(f"Error extracting text from PDF (PyPDF2): {e}")

    extract_text_from_pdf = extract_text_from_pdf_pypdf2

    @staticmethod
    def extract_text_from_docx(docx_path):
        current_app.logger.info(f"Attempting to extract text from DOCX: {docx_path}")
        try:
            doc = docx.Document(docx_path)
            text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            current_app.logger.info(f"Successfully extracted text from DOCX: {docx_path}. Length: {len(text)}")
            return text
        except Exception as e:
            current_app.logger.error(f"Error extracting text from DOCX {docx_path}: {e}", exc_info=True)
            raise DocumentProcessingError(f"Error extracting text from DOCX: {e}")

    @staticmethod
    def extract_text_from_txt(txt_path):
        current_app.logger.info(f"Attempting to extract text from TXT: {txt_path}")
        try:
            with open(txt_path, "r", encoding="utf-8") as f:
                text = f.read()
            current_app.logger.info(f"Successfully extracted text from TXT: {txt_path}. Length: {len(text)}")
            return text
        except Exception as e:
            current_app.logger.error(f"Error extracting text from TXT {txt_path}: {e}", exc_info=True)
            raise DocumentProcessingError(f"Error extracting text from TXT: {e}")

    @staticmethod
    def extract_text(file_path):
        _, ext = os.path.splitext(file_path)
        ext = ext.lower()
        current_app.logger.info(f"Extracting text from file: {file_path} (extension: {ext})")
        if ext == ".pdf":
            return TextExtractor.extract_text_from_pdf(file_path)
        elif ext == ".docx":
            return TextExtractor.extract_text_from_docx(file_path)
        elif ext == ".txt":
            return TextExtractor.extract_text_from_txt(file_path)
        else:
            current_app.logger.warning(f"Unsupported file type for text extraction: {ext} for file {file_path}")
            raise DocumentProcessingError(f"Unsupported file type for text extraction: {ext}")

class JDParser:
    _jd_parser_prompt_template_document = None

    @staticmethod
    def _load_prompt_template_document():
        if JDParser._jd_parser_prompt_template_document is None:
            JDParser._jd_parser_prompt_template_document = (
                "You are an expert HR analyst. Your task is to analyze the provided Job Description (JD) document "
                "and extract key requirements. Focus on identifying specific, actionable items for education, "
                "experience, and skills. Present these requirements in a structured JSON format. "
                "The JSON object must have three top-level keys: \"education\", \"experience\", and \"skills\". "
                "Each key should map to a list of strings, where each string is a distinct requirement. "
                "If a category has no specific items, provide an empty list for that category. "
                "Do not add any commentary or explanations outside of the JSON structure. "
                "Ensure your entire response is ONLY the JSON object itself, without any surrounding text or Markdown formatting.\n\n"
                "Required JSON Output Format Example:\n"
                "{\n"
                "  \"education\": [\"Bachelor\\u0027s degree in Computer Science or related field\"],\n"
                "  \"experience\": [\"5+ years of experience in software development\", \"Experience with Python and Django\"],\n"
                "  \"skills\": [\"Proficiency in JavaScript\", \"Strong problem-solving abilities\"]\n"
                "}"
            )
            current_app.logger.info("Using hardcoded JD parser LLM prompt template for DOCUMENT input.")
        return JDParser._jd_parser_prompt_template_document

    @staticmethod
    def _strip_markdown_code_block(text: str) -> str:
        text_to_log = text[:200] + "..." if len(text) > 200 else text
        current_app.logger.info(f"Attempting to strip Markdown. Input text (snippet): \n{text_to_log}\n")
        processed_text = text.strip()
        if processed_text.startswith("```json") and processed_text.endswith("```"):
            current_app.logger.info("Detected ```json ... ``` block.")
            processed_text = processed_text[len("```json"):-len("```")]
            processed_text = processed_text.strip()
        elif processed_text.startswith("```") and processed_text.endswith("```"):
            current_app.logger.info("Detected ``` ... ``` block.")
            processed_text = processed_text[len("```"):-len("```")]
            processed_text = processed_text.strip()
        else:
            current_app.logger.info("No Markdown code block detected by start/end checks. Text remains (after initial strip).")
        return processed_text

    @staticmethod
    def parse_jd(original_file_path: str, output_dir_for_pdf_conversion: str):
        current_app.logger.info(f"JDParser.parse_jd called for original file: {original_file_path}")
        try:
            converted_pdf_path = FileConverter.convert_to_pdf(original_file_path, output_dir_for_pdf_conversion)
            current_app.logger.info(f"Original file {original_file_path} converted to PDF: {converted_pdf_path}")

            if not converted_pdf_path or not os.path.exists(converted_pdf_path):
                current_app.logger.error(f"PDF conversion failed or file not found for {original_file_path}. Path: {converted_pdf_path}")
                raise DocumentProcessingError(f"Failed to convert or find PDF for {original_file_path}")

            parsed_requirements_dict = JDParser.parse_jd_from_document(converted_pdf_path)
            current_app.logger.info(f"Successfully parsed requirements from PDF {converted_pdf_path}")

            return {
                "requirements": parsed_requirements_dict,
                "processed_pdf_path": converted_pdf_path
            }
        except DocumentProcessingError as dpe:
            current_app.logger.error(f"DocumentProcessingError in JDParser.parse_jd for {original_file_path}: {dpe}", exc_info=True)
            raise
        except Exception as e:
            current_app.logger.error(f"Unexpected error in JDParser.parse_jd for {original_file_path}: {e}", exc_info=True)
            raise DocumentProcessingError(f"Unexpected error processing JD {original_file_path}: {e}")

    @staticmethod
    def parse_jd_from_document(pdf_file_path: str):
        current_app.logger.info(f"Attempting to parse JD from document: {pdf_file_path}")
        if not os.path.exists(pdf_file_path):
            current_app.logger.error(f"PDF file not found for JD parsing: {pdf_file_path}")
            raise DocumentProcessingError(f"PDF file not found: {pdf_file_path}")

        try:
            with open(pdf_file_path, "rb") as f:
                pdf_base64_data = base64.b64encode(f.read()).decode("utf-8")
            
            prompt_text = JDParser._load_prompt_template_document()
            llm_service = ClaudeService()
            current_app.logger.info("Calling LLM to parse JD requirements from PDF document.")
            
            messages_payload = [{
                "role": "user",
                "content": [
                    {
                        "type": "document",
                        "source": {
                            "type": "base64",
                            "media_type": "application/pdf",
                            "data": pdf_base64_data
                        }
                    },
                    {
                        "type": "text",
                        "text": prompt_text
                    }
                ]
            }]

            headers = {
                "x-api-key": llm_service.api_key,
                "anthropic-version": "2023-06-01",
                "content-type": "application/json"
            }
            payload = {
                "model": llm_service.model_name,
                "max_tokens": 2048,
                "messages": messages_payload
            }

            log_payload = json.loads(json.dumps(payload))
            if log_payload.get("messages") and log_payload["messages"][0].get("content") and \
               isinstance(log_payload["messages"][0]["content"], list) and \
               len(log_payload["messages"][0]["content"]) > 0 and \
               log_payload["messages"][0]["content"][0].get("source") and \
               log_payload["messages"][0]["content"][0]["source"].get("type") == "base64":
                log_payload["messages"][0]["content"][0]["source"]["data"] = log_payload["messages"][0]["content"][0]["source"]["data"][:100] + "... (truncated for logging)"
            current_app.logger.debug(f"Anthropic API Request Headers: {json.dumps(headers)}")
            current_app.logger.debug(f"Anthropic API Request Payload: {json.dumps(log_payload, indent=2)}")

            response = requests.post(llm_service.api_url, headers=headers, json=payload, timeout=180)
            current_app.logger.debug(f"Anthropic API Response Status Code: {response.status_code}")
            current_app.logger.debug(f"Anthropic API Response Headers: {json.dumps(dict(response.headers))}")
            current_app.logger.debug(f"Anthropic API Response Content: {response.text[:1000]}{'...' if len(response.text) > 1000 else ''}")
            response_data = response.json()
            if response_data.get("content") and isinstance(response_data["content"], list) and len(response_data["content"]) > 0:
                assistant_response_text = response_data["content"][0].get("text")
                if assistant_response_text:
                    current_app.logger.info(f"Raw LLM response for JD document parsing (first 500 chars): {assistant_response_text[:500]}")
                    text_to_parse_json = JDParser._strip_markdown_code_block(assistant_response_text)
                    try:
                        parsed_requirements = json.loads(text_to_parse_json)
                        if not all(k in parsed_requirements for k in ["education", "experience", "skills"]):
                            raise LLMAnalysisError("LLM response for JD parsing is missing required keys.")
                        current_app.logger.info("Successfully parsed JD requirements from document using LLM.")
                        return parsed_requirements
                    except json.JSONDecodeError as e:
                        raise DocumentProcessingError(f"Failed to parse LLM JSON response for JD document parsing: {e}")
                else:
                    raise DocumentProcessingError("LLM response content for JD document parsing is empty.")
            else:
                raise DocumentProcessingError(f"Unexpected LLM API response structure for JD document parsing. Full response: {response_data}")

        except LLMAnalysisError as e:
            current_app.logger.error(f"LLMAnalysisError during JD document parsing: {e}", exc_info=True)
            raise DocumentProcessingError(f"LLM Analysis Error during JD document parsing: {e}")
        except requests.exceptions.RequestException as e:
            current_app.logger.error(f"Error communicating with Anthropic API for JD document parsing: {e}", exc_info=True)
            if hasattr(e, "response") and e.response is not None:
                current_app.logger.error(f"Anthropic API Error Response Content: {e.response.text}")
            raise DocumentProcessingError(f"Error communicating with Anthropic API for JD document parsing: {e}")
        except Exception as e:
            current_app.logger.error(f"An unexpected error occurred during LLM JD document parsing: {e}", exc_info=True)
            raise DocumentProcessingError(f"An unexpected error occurred during LLM JD document parsing: {e}")

    @staticmethod
    def parse_jd_text_with_llm(jd_text):
        pass 

    @staticmethod
    def parse_jd_text_rule_based(jd_text):
        pass

